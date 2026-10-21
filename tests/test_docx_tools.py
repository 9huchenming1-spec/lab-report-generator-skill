"""Regression tests for template fidelity and fail-closed slot filling."""

import hashlib
import base64
from io import BytesIO
import importlib.util
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "skills/lab-report-generator/scripts"
spec = importlib.util.spec_from_file_location("docx_template", SCRIPTS / "docx_template.py")
tool = importlib.util.module_from_spec(spec)
spec.loader.exec_module(tool)


class TemplateTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        self.source = self.root / "模板.docx"
        self.plan = self.root / "plan.json"
        self.output = self.root / "report.docx"
        document = Document()
        paragraph = document.add_paragraph()
        paragraph.add_run("姓名：").bold = True
        paragraph.add_run("{{姓").italic = True
        paragraph.add_run("名}}")
        paragraph.add_run(" 保留文字").underline = True
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "答案"
        document.sections[0].header.paragraphs[0].text = "固定页眉"
        document.sections[0].footer.paragraphs[0].text = "固定页脚"
        field = document.add_paragraph("页码字段")
        run = OxmlElement("w:r")
        instruction = OxmlElement("w:instrText")
        instruction.text = " PAGE "
        run.append(instruction)
        field._p.append(run)
        image = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO6p7zgAAAAASUVORK5CYII=")
        document.add_picture(BytesIO(image))
        document.save(self.source)
        self.before = self.source.read_bytes()
        self.inventory = tool.inspect(self.source)

    def slot(self, text):
        return next(p for p in self.inventory["paragraphs"] if p["text"] == text and p["editable"])

    def write_plan(self, operations):
        self.plan.write_text(json.dumps({"template_sha256": self.inventory["template_sha256"],
                                         "operations": operations}), encoding="utf-8")

    def test_split_run_replace_and_empty_table_slot_preserve_package(self):
        name = self.slot("姓名：{{姓名}} 保留文字")
        empty = self.slot("")
        self.write_plan([
            {"part": name["part"], "path": name["path"], "expected": name["text"],
             "replace": {"old": "{{姓名}}", "new": "示例同学"}},
            {"part": empty["part"], "path": empty["path"], "expected": "",
             "text": "算法说明\n代码\t输出"},
        ])
        tool.fill(self.source, self.plan, self.output)
        document = Document(self.output)
        self.assertEqual(document.paragraphs[0].text, "姓名：示例同学 保留文字")
        self.assertTrue(document.paragraphs[0].runs[0].bold)
        self.assertTrue(document.paragraphs[0].runs[1].italic)
        self.assertTrue(document.paragraphs[0].runs[-1].underline)
        self.assertEqual(document.tables[0].cell(0, 1).text, "算法说明\n代码\t输出")
        self.assertEqual(self.source.read_bytes(), self.before)
        self.assertEqual(tool.compare(self.source, self.output)["changed"], ["word/document.xml"])
        self.assertEqual(tool.compare(self.source, self.output)["removed"], [])

    def test_wrong_original_text_leaves_no_output(self):
        slot = self.slot("答案")
        self.write_plan([{**{k: slot[k] for k in ("part", "path")}, "expected": "过期文本", "text": "新答案"}])
        with self.assertRaisesRegex(ValueError, "mismatch"):
            tool.fill(self.source, self.plan, self.output)
        self.assertFalse(self.output.exists())

    def test_stale_template_hash_is_rejected(self):
        self.write_plan([{}])
        plan = json.loads(self.plan.read_text())
        plan["template_sha256"] = "0" * 64
        self.plan.write_text(json.dumps(plan))
        with self.assertRaisesRegex(ValueError, "SHA-256"):
            tool.fill(self.source, self.plan, self.output)
        self.assertFalse(self.output.exists())

    def test_field_paragraph_is_rejected(self):
        field = next(p for p in self.inventory["paragraphs"] if p["text"] == "页码字段")
        self.assertFalse(field["editable"])
        self.write_plan([{**{k: field[k] for k in ("part", "path")}, "expected": field["text"], "text": "坏修改"}])
        with self.assertRaisesRegex(ValueError, "Unsafe"):
            tool.fill(self.source, self.plan, self.output)
        self.assertFalse(self.output.exists())

    def test_input_and_existing_output_cannot_be_overwritten(self):
        self.write_plan([{}])
        with self.assertRaises(ValueError):
            tool.fill(self.source, self.plan, self.source)
        self.output.write_text("existing", encoding="utf-8")
        with self.assertRaises(ValueError):
            tool.fill(self.source, self.plan, self.output)
        self.assertEqual(self.output.read_text(), "existing")
        self.assertEqual(self.source.read_bytes(), self.before)

    def test_pdf_output_is_not_accepted_as_docx(self):
        with self.assertRaises(ValueError):
            tool.fill(self.source, self.plan, self.root / "report.pdf")

    def test_invalid_conversion_creates_no_output(self):
        process = subprocess.run([sys.executable, str(SCRIPTS / "convert_document.py"), str(self.source),
                                  "--output", str(self.root / "report.doc")], capture_output=True)
        self.assertNotEqual(process.returncode, 0)
        self.assertFalse((self.root / "report.doc").exists())

    def test_missing_converter_does_not_claim_success(self):
        process = subprocess.run([sys.executable, str(SCRIPTS / "convert_document.py"), str(self.source),
            "--output", str(self.root / "report.pdf"), "--engine", "libreoffice",
            "--soffice", str(self.root / "missing")], capture_output=True)
        self.assertNotEqual(process.returncode, 0)
        self.assertFalse((self.root / "report.pdf").exists())

    def test_header_fill_changes_only_header_part(self):
        header = self.slot("固定页眉")
        self.write_plan([{"part": header["part"], "path": header["path"],
                          "expected": header["text"], "text": "已核对的页眉"}])
        tool.fill(self.source, self.plan, self.output)
        self.assertEqual(tool.compare(self.source, self.output)["changed"], [header["part"]])
        self.assertEqual(Document(self.output).sections[0].header.paragraphs[0].text, "已核对的页眉")

    def test_protected_template_is_rejected(self):
        document = Document(self.source)
        protection = OxmlElement("w:documentProtection")
        protection.set(qn("w:enforcement"), "1")
        document.settings.element.append(protection)
        protected = self.root / "protected.docx"
        document.save(protected)
        inventory = tool.inspect(protected)
        slot = next(p for p in inventory["paragraphs"] if p["text"] == "答案")
        self.plan.write_text(json.dumps({"template_sha256": inventory["template_sha256"],
            "operations": [{"part": slot["part"], "path": slot["path"], "expected": "答案", "text": "修改"}]}), encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "Protected"):
            tool.fill(protected, self.plan, self.output)
        self.assertFalse(self.output.exists())

    def test_ambiguous_replacement_is_rejected(self):
        slot = self.slot("姓名：{{姓名}} 保留文字")
        self.write_plan([{"part": slot["part"], "path": slot["path"], "expected": slot["text"],
                          "replace": {"old": "姓", "new": "名"}}])
        with self.assertRaisesRegex(ValueError, "exactly once"):
            tool.fill(self.source, self.plan, self.output)
        self.assertFalse(self.output.exists())


class PdfTests(unittest.TestCase):
    def test_blank_pdf_requires_visual_review(self):
        from pypdf import PdfWriter
        with tempfile.TemporaryDirectory() as directory:
            source, output = Path(directory) / "scan.pdf", Path(directory) / "scan.txt"
            writer = PdfWriter()
            writer.add_blank_page(width=300, height=300)
            with source.open("wb") as stream:
                writer.write(stream)
            process = subprocess.run([sys.executable, str(SCRIPTS / "extract_pdf.py"), str(source),
                                      "--output", str(output)], capture_output=True)
            self.assertEqual(process.returncode, 2)
            self.assertIn("NEEDS_VISUAL_REVIEW", output.read_text(encoding="utf-8"))


if __name__ == "__main__":
    unittest.main()
