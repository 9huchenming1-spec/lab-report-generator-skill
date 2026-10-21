"""Build a synthetic two-task example and fill it using the shipped template tool.

This is a reproducible demo fixture, not a general automatic report writer.
"""

import argparse
from datetime import datetime, timezone
import hashlib
import importlib.util
import inspect
import json
from pathlib import Path
import platform
import subprocess
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location("template_tool", ROOT / "skills/lab-report-generator/scripts/docx_template.py")
template_tool = importlib.util.module_from_spec(spec)
spec.loader.exec_module(template_tool)


def configure_font(style, name, size):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体" if name != "黑体" else "黑体")


def add_task(document, title, prefix):
    document.add_heading(title, level=1)
    table = document.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(2.5)
    table.columns[1].width = Cm(14.5)
    labels = [("设计思路", "DESIGN"), ("关键代码", "CODE"), ("运行结果", "RESULT"), ("结果分析", "ANALYSIS")]
    for row, (label, suffix) in zip(table.rows, labels):
        row.cells[0].width, row.cells[1].width = Cm(2.5), Cm(14.5)
        row.cells[0].text = label
        row.cells[1].text = "{{" + prefix + "_" + suffix + "}}"
        row.cells[0].paragraphs[0].runs[0].bold = True
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            properties = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right"):
                border = OxmlElement("w:" + edge)
                for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                    border.set(qn("w:" + key), value)
                borders.append(border)
            properties.append(borders)
            padding = OxmlElement("w:tcMar")
            for edge in ("top", "left", "bottom", "right"):
                margin = OxmlElement("w:" + edge)
                margin.set(qn("w:w"), "90")
                margin.set(qn("w:type"), "dxa")
                padding.append(margin)
            properties.append(padding)
        shade = OxmlElement("w:shd")
        shade.set(qn("w:fill"), "F2F2F2")
        row.cells[0]._tc.get_or_add_tcPr().append(shade)
        for paragraph in row.cells[1].paragraphs:
            if suffix in {"CODE", "RESULT"}:
                paragraph.style = document.styles["Code"]


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=ROOT / "examples/basic")
    args = parser.parse_args()
    output = args.output.resolve()
    output.mkdir(parents=True, exist_ok=True)
    names = ["template.docx", "report.docx", "fill-plan.json", "evidence.json", "run-results.txt"]
    if any((output / name).exists() for name in names):
        raise ValueError("Demo outputs already exist; choose a new output directory.")
    code = ROOT / "examples/basic/answer.py"
    process = subprocess.run([sys.executable, str(code)], capture_output=True, text=True, timeout=10, check=True)
    code_spec = importlib.util.spec_from_file_location("demo_answer", code)
    answer = importlib.util.module_from_spec(code_spec)
    code_spec.loader.exec_module(answer)
    original = [5, 2, 4, 2]
    assert answer.insertion_sort(original) == [2, 2, 4, 5] and original == [5, 2, 4, 2]
    assert answer.insertion_sort([]) == [] and answer.insertion_sort([7]) == [7]
    assert [answer.binary_search([1, 3, 5, 7, 9], v) for v in (7, 4)] == [3, -1]
    assert answer.binary_search([], 7) == -1
    (output / "run-results.txt").write_text(process.stdout, encoding="utf-8")
    document = Document()
    # The runtime's blank document may carry decorative Title/Subtitle borders.
    for border in list(document.styles.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
    document.core_properties.author = "Synthetic demonstration"
    document.core_properties.title = "排序与二分查找实验报告"
    document.core_properties.comments = "Public synthetic example; no real student data."
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(2)
    configure_font(document.styles["Normal"], "宋体", 10.5)
    document.styles["Normal"].paragraph_format.space_after = Pt(5)
    document.styles["Normal"].paragraph_format.line_spacing = 1.1
    configure_font(document.styles["Title"], "黑体", 20)
    configure_font(document.styles["Heading 1"], "黑体", 13)
    document.styles.add_style("Code", 1)
    configure_font(document.styles["Code"], "Consolas", 9)
    document.styles["Code"].paragraph_format.space_after = Pt(4)
    document.styles["Code"].paragraph_format.line_spacing = 1.0
    section.header.paragraphs[0].text = "程序设计实验报告  公开演示材料"
    footer = section.footer.paragraphs[0]
    footer.text = "演示模板  第 "
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.add_run(" 页")
    document.add_paragraph("排序与二分查找实验报告", style="Title")
    metadata = document.add_paragraph("姓名：")
    metadata.add_run("{{姓")
    metadata.add_run("名}}")
    metadata.add_run("    班级：{{班级}}    学号：{{学号}}")
    document.add_paragraph("实验日期：{{日期}}")
    document.add_heading("实验目的", level=1)
    document.add_paragraph("{{目的}}")
    add_task(document, "任务1 插入排序", "T1")
    document.add_page_break()
    add_task(document, "任务2 二分查找", "T2")
    document.add_heading("实验总结", level=1)
    document.add_paragraph("{{总结}}")
    document.save(output / "template.docx")
    mapping = template_tool.inspect(output / "template.docx")
    lines = process.stdout.strip().splitlines()
    values = {
        "姓名：{{姓名}}    班级：{{班级}}    学号：{{学号}}": "姓名：示例同学    班级：演示班    学号：DEMO-001",
        "实验日期：{{日期}}": "实验日期：2026-10-21",
        "{{目的}}": "实现插入排序与二分查找，分析循环和边界更新，使用正常及边界输入核对算法行为。",
        "{{T1_DESIGN}}": "先复制输入列表，再将下标 index 处的元素插入已经有序的前缀。大于 current 的元素向右移动，空出的最终位置写入 current；因此原列表保持不变。",
        "{{T1_CODE}}": inspect.getsource(answer.insertion_sort).strip(),
        "{{T1_RESULT}}": "本次实际运行 Python 演示程序，退出码为 0。\n" + "\n".join(lines[:3]),
        "{{T1_ANALYSIS}}": "普通输入得到 [2, 2, 4, 5]，重复值被保留；原列表仍为 [5, 2, 4, 2]。空列表和单元素均返回正确结果。由循环分析，最坏/平均时间复杂度为 O(n²)，已排序时为 O(n)；列表副本使额外空间为 O(n)。",
        "{{T2_DESIGN}}": "使用闭区间 [left, right]。若中点小于目标，令 left = middle + 1；否则令 right = middle - 1。每次丢弃不含目标的一半，区间为空时返回 -1。",
        "{{T2_CODE}}": inspect.getsource(answer.binary_search).strip(),
        "{{T2_RESULT}}": "本次实际运行输出如下；原始记录见公开演示的 run-results.txt。\n" + "\n".join(lines[3:]),
        "{{T2_ANALYSIS}}": "目标 7 位于下标 3；不存在的目标 4 和空列表均返回 -1。代码以升序输入为前提，不检查或排序输入。由区间每轮减半可得时间复杂度 O(log n)，使用常数个变量，额外空间 O(1)。",
        "{{总结}}": "两个函数在本次指定的正常与边界输入下得到预期结果。插入排序适合解释有序前缀的维护；二分查找依赖输入已排序。本次没有测量耗时，也未验证全部可能输入，不能据此推断性能实测或穷尽正确性。",
    }
    operations = []
    for paragraph in mapping["paragraphs"]:
        if paragraph["text"] in values:
            operations.append({"part": paragraph["part"], "path": paragraph["path"],
                               "expected": paragraph["text"], "text": values[paragraph["text"]]})
    assert len(operations) == len(values)
    plan = {"template_sha256": mapping["template_sha256"], "operations": operations}
    (output / "fill-plan.json").write_text(json.dumps(plan, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    template_tool.fill(output / "template.docx", output / "fill-plan.json", output / "report.docx")
    evidence = {"kind": "executed", "created_utc": datetime.now(timezone.utc).isoformat(),
        "command": ["python", "examples/basic/answer.py"], "python": platform.python_version(),
        "platform": platform.system(), "source_sha256": hashlib.sha256(code.read_bytes()).hexdigest(),
        "exit_code": process.returncode, "stdout": process.stdout, "stderr": process.stderr,
        "requirements": [
            {"id": "任务1", "code": "answer.py:insertion_sort", "slots": "任务1四个表格栏目", "evidence": "run-results.txt前3行", "status": "complete"},
            {"id": "任务2", "code": "answer.py:binary_search", "slots": "任务2四个表格栏目", "evidence": "run-results.txt后3行", "status": "complete"}],
        "package_comparison": template_tool.compare(output / "template.docx", output / "report.docx")}
    (output / "evidence.json").write_text(json.dumps(evidence, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"output": str(output), "operations": len(operations), "executed_cases": 6}, ensure_ascii=False))


if __name__ == "__main__":
    main()
